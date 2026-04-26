import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Title
title = doc.add_heading('Unibuddy Project Report', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

subtitle = doc.add_paragraph('A Comprehensive Overview of the SRMAP Attendance Intelligence Platform')
subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
subtitle.style.font.italic = True

doc.add_paragraph('\n')

# 1. Introduction
doc.add_heading('1. Introduction', level=1)
doc.add_paragraph(
    'Unibuddy is a modern, intelligent web application designed specifically for students of '
    'SRM University AP. It automates the process of tracking academic progress by integrating '
    'live data from the SRMAP student portal. The platform provides a seamless dashboard for students '
    'to monitor their attendance, check their daily timetable, and receive smart analytics such as '
    '"bunk estimation" (calculating how many classes a student can safely miss while maintaining '
    'the required 75% attendance criteria).'
)

# 2. Core Features
doc.add_heading('2. Core Features', level=1)
features = [
    ('Automated Portal Integration', 'Eliminates the need for manual data entry by logging into the SRM AP portal on behalf of the user.'),
    ('AI-Powered Captcha Solving', 'Incorporates an automated captcha solver using a Convolutional Recurrent Neural Network (CRNN) to bypass the university portal\'s captcha transparently.'),
    ('Attendance Analytics & Bunk Estimator', 'Provides deep insights into attendance percentages. It clearly shows "safe" and "danger" zones, calculating exactly how many classes a student must attend or can afford to miss.'),
    ('Interactive Timetable & Academic Calendar', 'Visualizes the student\'s daily schedule symmetrically, alongside holiday and academic calendar data, in a highly responsive "Liquid Glass" gooey UI interface.'),
    ('Session Management', 'Implements a "heartbeat" background mechanism to keep the university portal session active and prevent automatic logouts.'),
    ('Theme-Aware Aesthetics', 'Features a beautiful dynamic interface with a Red/Black dark mode and a high-contrast Black/White light mode.')
]

for title, desc in features:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(title + ': ').bold = True
    p.add_run(desc)

# 3. Architecture & Tech Stack
doc.add_heading('3. Architecture & Tech Stack', level=1)
doc.add_paragraph('The project is architected as a modern hybrid application with three main components:')

doc.add_heading('A. Frontend (User Interface)', level=2)
p_front = doc.add_paragraph()
p_front.add_run('Technologies: ').bold = True
p_front.add_run('React, Vite, TypeScript, Tailwind CSS\n')
p_front.add_run(
    'The frontend is built for extreme responsiveness and visual excellence, featuring glassmorphism '
    'and modern UI components. It securely stores session identifiers and communicates with the backend APIs.'
)

doc.add_heading('B. Main Backend (Data & API Layer)', level=2)
p_back = doc.add_paragraph()
p_back.add_run('Technologies: ').bold = True
p_back.add_run('Node.js, Express, TypeScript, SQLite\n')
p_back.add_run(
    'Acts as the orchestrator. It handles user authentication, portal scraping, and intermediate '
    'database storage. The backend uses Cheerio to parse HTML from the university portal and map it '
    'to structured JSON data (Profile, Attendance, Timetable).'
)

doc.add_heading('C. Captcha Solver Service', level=2)
p_ai = doc.add_paragraph()
p_ai.add_run('Technologies: ').bold = True
p_ai.add_run('Python, FastAPI, ONNX Runtime, PyTorch (Vision)\n')
p_ai.add_run(
    'A specialized microservice that receives base64 captcha images from the Node.js scraper, '
    'processes them using a trained CRNN ONNX model, and returns the solved text.'
)

# 4. System Workflow
doc.add_heading('4. System Workflow', level=1)
workflow = [
    ('User Login', 'The user enters their register number and password on the sleek frontend interface.'),
    ('Portal Authentication', 'The Node.js backend initiates a session with the SRM AP portal, downloads the login captcha, and sends it to the Python Captcha Solver.'),
    ('Machine Learning Inference', 'The fastAPI service cleans the captcha image, runs it through the CRNN model, and decodes the characters.'),
    ('Data Scraping', 'The Node.js backend logs into the portal using the credentials and the solved captcha. Once inside, it posts requests to fetch attendance, timetable, and profile data in parallel.'),
    ('Content Delivery', 'The parsed data is sent to the frontend, where it is beautifully visualized on the dashboard. A background heartbeat keeps the portal session alive.')
]

for i, (title, desc) in enumerate(workflow):
    p = doc.add_paragraph(style='List Number')
    p.add_run(title + ': ').bold = True
    p.add_run(desc)

# 5. Value Proposition
doc.add_heading('5. Value Proposition', level=1)
doc.add_paragraph(
    'Unibuddy solves a core pain point for students—the clunky, slow, and manual process of checking '
    'university portals. By automating the data retrieval and enhancing it with predictive analytics '
    '(like the bunk estimator), Unibuddy transforms raw university data into actionable academic intelligence.'
)

doc.save("Unibuddy_Project_Report.docx")
print("Report generated successfully as Unibuddy_Project_Report.docx")
